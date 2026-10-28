"""Export one student per row for SPSS analysis of the five Table 6 fields."""
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill

from generate_figures_r1_12 import load_verified, SCORE_COLUMNS

ROOT = Path(__file__).resolve().parent


def main():
    d, sources = load_verified()
    median = float(d.Effort_Total.median())
    export = pd.DataFrame({
        'Student_Index': d.Index.astype(int),
        'Student_ID': d.Id.astype(str),
        'Composite_Effort': d.Effort_Total,
        'Comment_Density': d.Comment_density,
        'R1_Start': d['R1.Start'],
        'Round2_Retention': d.Effort_R2 / d.Effort_R1,
        'Mean_Assignment_Score': d.Mean_assignment_score_6,
        'Cluster_Group': d.Cluster.astype(int),
        'Cluster_Label': d.Cluster.map({1: 'Extended Engagers', 2: 'Brief Participants'}),
        'Median_Group': np.where(d.Effort_Total > median, 1, 2),
        'Median_Label': np.where(d.Effort_Total > median, 'High effort', 'Low effort'),
    })
    definitions = [
        ('Student_Index', 'Numeric', 'Nominal', 'Original Index in Merge-Data.xlsx', 'Student identification only'),
        ('Student_ID', 'String', 'Nominal', 'Original Id in Merge-Data.xlsx', 'Student identification only'),
        ('Composite_Effort', 'Numeric', 'Scale', 'Sum of six R1/R2 C1/C2/C3 submission-count columns', 'Raw count'),
        ('Comment_Density', 'Numeric', 'Scale', '(R1_MF_Count + R2_MF_Count) / Composite_Effort', 'Ratio, not percentage; retain full precision'),
        ('R1_Start', 'Numeric', 'Scale', 'R1.Start from Merge-Data.xlsx', 'Raw recorded time units, no conversion'),
        ('Round2_Retention', 'Numeric', 'Scale', 'Effort_R2 / Effort_R1', 'Count ratio; can exceed 1'),
        ('Mean_Assignment_Score', 'Numeric', 'Scale', 'Mean of ' + ', '.join(SCORE_COLUMNS), 'Six-column definition in primary Table 6; not stored seven-column Mean_HW'),
        ('Cluster_Group', 'Numeric', 'Nominal', 'Saved Cluster; 1=Extended Engagers, 2=Brief Participants', 'n=41 and n=32; no reclustering'),
        ('Cluster_Label', 'String', 'Nominal', 'Human-readable label for Cluster_Group', 'Use numeric Cluster_Group for grouping in SPSS'),
        ('Median_Group', 'Numeric', 'Nominal', '1=High effort (>364), 2=Low effort (<=364)', 'n=36 and n=37; equality at median assigned Low effort'),
        ('Median_Label', 'String', 'Nominal', 'Human-readable label for Median_Group', 'Use numeric Median_Group for grouping in SPSS'),
    ]
    codebook = pd.DataFrame(definitions, columns=['Variable', 'Type', 'SPSS_Measure', 'Definition_or_value_labels', 'Notes'])
    counts = pd.concat([
        export.groupby(['Cluster_Group', 'Cluster_Label']).size().rename('N').reset_index().rename(columns={'Cluster_Group': 'Code', 'Cluster_Label': 'Label'}).assign(Grouping='Cluster'),
        export.groupby(['Median_Group', 'Median_Label']).size().rename('N').reset_index().rename(columns={'Median_Group': 'Code', 'Median_Label': 'Label'}).assign(Grouping='Median split'),
    ], ignore_index=True)[['Grouping', 'Code', 'Label', 'N']]
    notes = pd.DataFrame([
        ('Primary source', str(sources[0])),
        ('Verification source', str(sources[1])),
        ('Verification', 'Student IDs, saved cluster labels and behavior/count values match across both workbooks'),
        ('Students', '73; one row per student; no missing analysis values'),
        ('Score definition', 'Mean of HW1, HW3, HW4, HW5, HW6, HW7; HW2 equals HW1 for every student and is counted once'),
        ('Original stored Mean_HW', 'Not exported as the main score: it averages seven columns including identical HW1/HW2 recordings'),
        ('Median threshold', '364; High if >364; Low if <=364'),
        ('SPSS import', 'Import the Student_Data worksheet and read variable names from its first row'),
        ('Measurement levels', 'Set the five analysis variables to Scale; identifiers and both group variables to Nominal'),
        ('Adjusted regression coding', 'To reproduce Table 6 b, use a 0/1 indicator: 1 for Cluster_Group=1, 0 for Cluster_Group=2; also include Composite_Effort as a covariate'),
        ('Precision', 'Numeric values are stored at full precision, not rounded to displayed Table 6 values'),
        ('Interpretation', 'Behavioral features contributed to clustering; their comparisons are descriptive, not independent validation of behavioral types'),
    ], columns=['Item', 'Details'])
    fields = ['Composite_Effort', 'Comment_Density', 'R1_Start', 'Round2_Retention', 'Mean_Assignment_Score']
    assert len(export) == 73 and export.Student_ID.is_unique
    assert np.isfinite(export[fields].to_numpy()).all()
    assert export.Cluster_Group.value_counts().to_dict() == {1: 41, 2: 32}
    assert export.Median_Group.value_counts().to_dict() == {2: 37, 1: 36}
    path = ROOT / 'Table6_results' / 'Table6_Student_Data_SPSS.xlsx'
    path.parent.mkdir(exist_ok=True)
    with pd.ExcelWriter(path, engine='openpyxl') as writer:
        for name, frame in [('Student_Data', export), ('Codebook', codebook), ('Group_Counts', counts), ('Source_and_Notes', notes)]:
            frame.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.sheets[name]
            sheet.freeze_panes = 'A2'
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill('solid', fgColor='DCE6F1')
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = min(85, max(16, max(len(str(c.value or '')) for c in column) + 2))
        sheet = writer.sheets['Student_Data']
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if sheet.cell(1, cell.column).value in fields:
                    cell.number_format = '0.00000000'
            row[1].number_format = '@'
    print(path)
    print(counts.to_string(index=False))


if __name__ == '__main__':
    main()
